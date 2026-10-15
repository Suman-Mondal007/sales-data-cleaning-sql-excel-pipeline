"""
run_pipeline.py  --  one command runs the whole workflow:

    python run_pipeline.py

1. Extract   : read 3 raw CSVs (orders, customers, products)
2. Clean     : fix types, formats, labels, duplicates, missing values
3. Validate  : rule checks; failing rows are quarantined with a reason
4. Load      : write clean tables to SQLite (db/sales.db)
5. Analyse   : run the SQL queries in sql/queries.sql
6. Report    : data-quality report, PNG charts and an Excel dashboard
"""
import json
import re
import sqlite3
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).parent
RAW, CLEAN, REJ = ROOT / "data/raw", ROOT / "data/clean", ROOT / "data/rejected"
DB_PATH, REPORTS, DOCS = ROOT / "db/sales.db", ROOT / "reports", ROOT / "docs"
SQL_OUT = REPORTS / "sql_results"

# ---- business rules (documented assumptions) ----
MIN_DATE, MAX_DATE = pd.Timestamp("2024-01-01"), pd.Timestamp("2025-12-31")
QTY_MIN, QTY_MAX = 1, 50
DISC_MIN, DISC_MAX = 0.0, 0.5
STATUS_OK = {"Delivered", "Cancelled", "Returned", "Pending"}
PAY_OK = {"UPI", "Credit Card", "Debit Card", "Net Banking", "Cash on Delivery"}
EMAIL_RE = re.compile(r"^[\w.+-]+@[\w-]+\.[\w.]+$")

CITY_ALIAS = {"calcutta": "Kolkata", "kolkatta": "Kolkata", "bangalore": "Bengaluru",
              "ahmadabad": "Ahmedabad"}
REGION = {"Kolkata": "East", "Mumbai": "West", "Pune": "West", "Ahmedabad": "West",
          "Delhi": "North", "Bengaluru": "South", "Chennai": "South", "Hyderabad": "South"}
CATEGORY_ALIAS = {"home and kitchen": "Home & Kitchen"}
PAY_ALIAS = {"upi": "UPI", "cc": "Credit Card", "credit card": "Credit Card", "dc": "Debit Card",
             "debit card": "Debit Card", "netbanking": "Net Banking", "net banking": "Net Banking",
             "cod": "Cash on Delivery", "cash on delivery": "Cash on Delivery"}

metrics: dict = {}


# ------------------------------------------------------------------ helpers
def title_label(series: pd.Series, alias: dict) -> pd.Series:
    key = series.astype("string").str.strip().str.lower()
    return key.map(alias).fillna(key.str.title())


def parse_date(x):
    if pd.isna(x):
        return pd.NaT
    x = str(x).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%b %d, %Y"):
        try:
            return pd.Timestamp(datetime.strptime(x, fmt))
        except ValueError:
            continue
    return pd.NaT


def parse_money(s: pd.Series) -> pd.Series:
    cleaned = s.astype("string").str.replace(r"(?i)rs\.?|\u20b9|,|\s", "", regex=True)
    return pd.to_numeric(cleaned, errors="coerce")


def count_label_variants(raw: pd.Series, clean: pd.Series) -> tuple[int, int]:
    return int(raw.astype("string").nunique()), int(clean.nunique())


# ------------------------------------------------------------------ 1. extract
def extract():
    kw = dict(dtype=str, keep_default_na=True)
    orders = pd.read_csv(RAW / "orders_raw.csv", **kw)
    customers = pd.read_csv(RAW / "customers_raw.csv", **kw)
    products = pd.read_csv(RAW / "products_raw.csv", **kw)
    metrics["raw_rows"] = {"orders": len(orders), "customers": len(customers), "products": len(products)}
    return orders, customers, products


# ------------------------------------------------------------------ 2. clean
def clean_products(df: pd.DataFrame, orders_raw: pd.DataFrame) -> pd.DataFrame:
    p = df.copy()
    p["product_id"] = p["product_id"].str.strip().str.upper()
    raw_cat = p["category"].copy()
    p["category"] = title_label(p["category"], CATEGORY_ALIAS)
    p["category"] = p["category"].replace({"Home & Kitchen": "Home & Kitchen"})
    p["list_price"] = pd.to_numeric(p["list_price"], errors="coerce")
    p = p.drop_duplicates("product_id")

    # impute missing master price with the most common price seen in orders
    missing = p["list_price"].isna()
    metrics["products_missing_price"] = int(missing.sum())
    if missing.any():
        o = orders_raw[["product_id", "unit_price"]].copy()
        o["product_id"] = o["product_id"].str.strip().str.upper()
        o["unit_price"] = parse_money(o["unit_price"])
        mode_price = o.dropna().groupby("product_id")["unit_price"].agg(lambda x: x.mode().iloc[0])
        p.loc[missing, "list_price"] = p.loc[missing, "product_id"].map(mode_price)
    metrics["products_still_missing_price"] = int(p["list_price"].isna().sum())
    metrics["category_variants"] = count_label_variants(raw_cat, p["category"])
    return p


def clean_customers(df: pd.DataFrame) -> pd.DataFrame:
    c = df.copy()
    c["customer_id"] = c["customer_id"].str.strip().str.upper()
    metrics["customers_duplicates_removed"] = int(c.duplicated("customer_id").sum())
    c = c.drop_duplicates("customer_id", keep="first")

    raw_city = c["city"].copy()
    c["city"] = title_label(c["city"], CITY_ALIAS)
    metrics["city_variants"] = count_label_variants(raw_city, c["city"])
    c["region"] = c["city"].map(REGION)

    metrics["emails_missing"] = int(c["email"].isna().sum())
    valid = c["email"].fillna("").str.strip().apply(lambda e: bool(EMAIL_RE.match(e)))
    metrics["emails_invalid"] = int((~valid & c["email"].notna()).sum())
    c.loc[~valid, "email"] = np.nan
    c["signup_date"] = pd.to_datetime(c["signup_date"], errors="coerce").dt.strftime("%Y-%m-%d")
    return c.reset_index(drop=True)


def clean_orders(raw: pd.DataFrame, customers: pd.DataFrame, products: pd.DataFrame):
    o = raw.copy()
    for col in ("order_id", "customer_id", "product_id"):
        o[col] = o[col].str.strip().str.upper()

    # duplicates (only visible after normalising keys)
    exact_before = int(raw.duplicated().sum())
    metrics["orders_exact_duplicates_before_normalising"] = exact_before
    metrics["orders_duplicates_removed"] = int(o.duplicated("order_id").sum())
    o = o.drop_duplicates("order_id", keep="first").copy()

    # labels
    raw_status, raw_pay = o["status"].copy(), o["payment_method"].copy()
    o["status"] = title_label(o["status"], {})
    o["payment_method"] = title_label(o["payment_method"], PAY_ALIAS)
    metrics["status_variants"] = count_label_variants(raw_status, o["status"])
    metrics["payment_variants"] = count_label_variants(raw_pay, o["payment_method"])

    # types
    raw_dates = o["order_date"].copy()
    o["order_date"] = raw_dates.apply(parse_date)
    metrics["date_formats_in_raw"] = 3
    o["quantity"] = pd.to_numeric(o["quantity"], errors="coerce")
    o["unit_price"] = parse_money(o["unit_price"])
    metrics["price_values_with_currency_symbols"] = int(
        raw["unit_price"].astype("string").str.contains(r"[^\d.]", na=False).sum())

    d = o["discount"].astype("string").str.strip()
    pct = d.str.endswith("%", na=False)
    o["discount"] = pd.to_numeric(d.str.rstrip("%"), errors="coerce")
    o.loc[pct, "discount"] = o.loc[pct, "discount"] / 100.0
    metrics["discount_missing_filled_with_0"] = int(o["discount"].isna().sum())
    o["discount"] = o["discount"].fillna(0.0)

    # impute missing price from product master
    price_missing = o["unit_price"].isna()
    metrics["unit_price_missing_filled_from_master"] = int(price_missing.sum())
    master = products.set_index("product_id")["list_price"]
    o.loc[price_missing, "unit_price"] = o.loc[price_missing, "product_id"].map(master)

    # ---- validation: flag failing rows with a reason ----
    o["reject_reason"] = pd.Series([None] * len(o), index=o.index, dtype="object")

    def flag(mask, reason):
        o.loc[mask & o["reject_reason"].isna(), "reject_reason"] = reason

    flag(o["order_date"].isna(), "invalid_or_unparseable_date")
    flag((o["order_date"] < MIN_DATE) | (o["order_date"] > MAX_DATE), "date_out_of_range")
    flag(o["quantity"].isna(), "missing_quantity")
    flag((o["quantity"] < QTY_MIN) | (o["quantity"] > QTY_MAX), "quantity_out_of_range")
    flag(~o["product_id"].isin(products["product_id"]), "unknown_product_id")
    flag(~o["customer_id"].isin(customers["customer_id"]), "unknown_customer_id")
    flag(o["unit_price"].isna() | (o["unit_price"] <= 0), "missing_or_invalid_price")
    flag((o["discount"] < DISC_MIN) | (o["discount"] > DISC_MAX), "discount_out_of_range")
    flag(~o["status"].isin(STATUS_OK), "invalid_status")
    flag(~o["payment_method"].isin(PAY_OK), "invalid_payment_method")

    rejected = o[o["reject_reason"].notna()].copy()
    good = o[o["reject_reason"].isna()].drop(columns="reject_reason").copy()
    good["quantity"] = good["quantity"].astype(int)
    good["order_value"] = (good["quantity"] * good["unit_price"] * (1 - good["discount"])).round(2)
    good["order_date"] = good["order_date"].dt.strftime("%Y-%m-%d")
    metrics["orders_rejected_by_reason"] = rejected["reject_reason"].value_counts().to_dict()
    return good.reset_index(drop=True), rejected


# ------------------------------------------------------------------ 3. validate
def validate(orders, customers, products) -> list[dict]:
    checks = []

    def check(name, ok):
        checks.append({"check": name, "result": "PASS" if bool(ok) else "FAIL"})

    check("order_id is unique", orders["order_id"].is_unique)
    check("customer_id is unique", customers["customer_id"].is_unique)
    check("product_id is unique", products["product_id"].is_unique)
    check("no nulls in key order columns",
          orders[["order_id", "customer_id", "product_id", "order_date", "quantity", "unit_price",
                  "status", "payment_method"]].notna().all().all())
    check("quantity within 1-50", orders["quantity"].between(QTY_MIN, QTY_MAX).all())
    check("unit_price > 0", (orders["unit_price"] > 0).all())
    check("discount within 0-50%", orders["discount"].between(DISC_MIN, DISC_MAX).all())
    check("every order has a valid customer", orders["customer_id"].isin(customers["customer_id"]).all())
    check("every order has a valid product", orders["product_id"].isin(products["product_id"]).all())
    check("order dates within 2024-01-01..2025-12-31",
          pd.to_datetime(orders["order_date"]).between(MIN_DATE, MAX_DATE).all())
    check("order_value = qty x price x (1 - discount)",
          np.allclose(orders["order_value"],
                      orders["quantity"] * orders["unit_price"] * (1 - orders["discount"]), atol=0.01))
    check("status values are standardised", set(orders["status"]) <= STATUS_OK)
    check("payment methods are standardised", set(orders["payment_method"]) <= PAY_OK)
    check("all cities mapped to a region", customers["region"].notna().all())
    return checks


# ------------------------------------------------------------------ 4. load
SCHEMA = """
DROP TABLE IF EXISTS orders; DROP TABLE IF EXISTS customers; DROP TABLE IF EXISTS products;
CREATE TABLE products  (product_id TEXT PRIMARY KEY, product_name TEXT, category TEXT, list_price REAL);
CREATE TABLE customers (customer_id TEXT PRIMARY KEY, customer_name TEXT, email TEXT, city TEXT,
                        region TEXT, signup_date TEXT);
CREATE TABLE orders (order_id TEXT PRIMARY KEY, customer_id TEXT NOT NULL REFERENCES customers(customer_id),
                     product_id TEXT NOT NULL REFERENCES products(product_id), order_date TEXT NOT NULL,
                     quantity INTEGER, unit_price REAL, discount REAL, status TEXT, payment_method TEXT,
                     order_value REAL);
CREATE INDEX idx_orders_customer ON orders(customer_id);
CREATE INDEX idx_orders_product  ON orders(product_id);
CREATE INDEX idx_orders_date     ON orders(order_date);
"""


def load_sqlite(orders, customers, products):
    DB_PATH.parent.mkdir(exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    con.executescript(SCHEMA)
    products.to_sql("products", con, if_exists="append", index=False)
    customers.to_sql("customers", con, if_exists="append", index=False)
    orders.to_sql("orders", con, if_exists="append", index=False)
    con.commit()
    return con


# ------------------------------------------------------------------ 5. analyse
def run_sql(con) -> dict:
    SQL_OUT.mkdir(parents=True, exist_ok=True)
    text = (ROOT / "sql/queries.sql").read_text()
    blocks = re.split(r"^-- name:\s*(\w+)\s*$", text, flags=re.M)[1:]
    results = {}
    for name, body in zip(blocks[0::2], blocks[1::2]):
        df = pd.read_sql_query(body.strip().rstrip(";"), con)
        df.to_csv(SQL_OUT / f"{name}.csv", index=False)
        results[name] = df
    return results


# ------------------------------------------------------------------ 6. report
def make_flat(con) -> pd.DataFrame:
    q = """SELECT o.order_id, o.order_date, substr(o.order_date,1,7) AS month, o.customer_id,
                  c.city, c.region, o.product_id, p.category, o.quantity, o.unit_price,
                  o.discount, o.order_value, o.status, o.payment_method
           FROM orders o JOIN customers c ON c.customer_id=o.customer_id
                         JOIN products  p ON p.product_id=o.product_id
           ORDER BY o.order_date, o.order_id"""
    return pd.read_sql_query(q, con)


def write_quality_report(checks):
    m = metrics
    rej = m["orders_rejected_by_reason"]
    total_rej = sum(rej.values())
    raw_o, clean_o = m["raw_rows"]["orders"], m["clean_rows"]["orders"]
    lines = [
        "# Data Quality Report", "",
        f"_Generated {datetime.now():%Y-%m-%d %H:%M}_", "",
        "## Row counts (before -> after)", "",
        "| Table | Raw rows | Clean rows |", "|---|---:|---:|",
        f"| orders | {raw_o:,} | {clean_o:,} |",
        f"| customers | {m['raw_rows']['customers']:,} | {m['clean_rows']['customers']:,} |",
        f"| products | {m['raw_rows']['products']:,} | {m['clean_rows']['products']:,} |", "",
        "## Issues found and how they were handled", "",
        "| Issue | Count | Action |", "|---|---:|---|",
        f"| Duplicate order rows (after trimming IDs) | {m['orders_duplicates_removed']:,} | Removed, kept first |",
        f"| Duplicate customer rows | {m['customers_duplicates_removed']:,} | Removed, kept first |",
        f"| Status label variants | {m['status_variants'][0]} -> {m['status_variants'][1]} | Standardised |",
        f"| Payment label variants | {m['payment_variants'][0]} -> {m['payment_variants'][1]} | Standardised |",
        f"| City label variants | {m['city_variants'][0]} -> {m['city_variants'][1]} | Standardised, mapped to region |",
        f"| Category label variants | {m['category_variants'][0]} -> {m['category_variants'][1]} | Standardised |",
        f"| Prices with currency symbols / commas | {m['price_values_with_currency_symbols']:,} | Parsed to numbers |",
        f"| Missing unit prices | {m['unit_price_missing_filled_from_master']:,} | Filled from product master |",
        f"| Missing product master prices | {m['products_missing_price']} | Filled with most common order price |",
        f"| Missing discounts | {m['discount_missing_filled_with_0']:,} | Assumed 0% |",
        f"| Missing customer emails | {m['emails_missing']} | Left blank |",
        f"| Invalid customer emails | {m['emails_invalid']} | Set to blank |",
        f"| Mixed date formats | {m['date_formats_in_raw']} formats | Parsed to ISO (YYYY-MM-DD) |", "",
        f"## Rows quarantined ({total_rej:,} orders -> `data/rejected/orders_rejected.csv`)", "",
        "| Reason | Rows |", "|---|---:|",
    ]
    lines += [f"| {k} | {v:,} |" for k, v in rej.items()]
    lines += ["", "## Post-cleaning validation", "", "| Check | Result |", "|---|---|"]
    lines += [f"| {c['check']} | {c['result']} |" for c in checks]
    passed = sum(c["result"] == "PASS" for c in checks)
    lines += ["", f"**{passed}/{len(checks)} checks passed.**", "",
              "## Assumptions", "",
              "- Dates containing `/` are read as DD/MM/YYYY (Indian convention).",
              "- Net revenue counts only orders with status `Delivered`.",
              "- A missing discount means no discount (0%).",
              f"- Quantity must be {QTY_MIN}-{QTY_MAX}; discount must be {DISC_MIN:.0%}-{DISC_MAX:.0%}; "
              f"dates must fall in {MIN_DATE:%Y-%m-%d} to {MAX_DATE:%Y-%m-%d}."]
    (REPORTS / "data_quality_report.md").write_text("\n".join(lines))
    (REPORTS / "metrics.json").write_text(json.dumps(m, indent=2, default=str))


def make_charts(res, flat):
    DOCS.mkdir(exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "axes.spines.top": False, "axes.spines.right": False})

    m = res["q02_monthly_revenue_mom"]
    fig, ax = plt.subplots(figsize=(10, 4.2))
    ax.plot(m["month"], m["revenue"] / 1e6, marker="o", color="#1F3864", linewidth=2)
    ax.set_title("Monthly net revenue (delivered orders)", loc="left", fontweight="bold")
    ax.set_ylabel("Revenue (INR million)")
    ax.tick_params(axis="x", rotation=60)
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout(); fig.savefig(DOCS / "monthly_revenue.png", dpi=150); plt.close(fig)

    c = res["q03_revenue_by_category"]
    fig, ax = plt.subplots(figsize=(8, 4.2))
    ax.barh(c["category"][::-1], c["revenue"][::-1] / 1e6, color="#2E75B6")
    for y, (v, s) in enumerate(zip(c["revenue"][::-1] / 1e6, c["share_pct"][::-1])):
        ax.text(v, y, f" {v:.1f}M ({s:.0f}%)", va="center", fontsize=9)
    ax.set_title("Net revenue by category", loc="left", fontweight="bold")
    ax.set_xlabel("Revenue (INR million)")
    ax.set_xlim(0, c["revenue"].max() / 1e6 * 1.25)
    fig.tight_layout(); fig.savefig(DOCS / "category_revenue.png", dpi=150); plt.close(fig)

    rc = res["q04_revenue_region_category"].pivot(index="region", columns="category", values="revenue") / 1e6
    fig, ax = plt.subplots(figsize=(9, 3.8))
    im = ax.imshow(rc.values, cmap="Blues", aspect="auto")
    ax.set_xticks(range(len(rc.columns)), rc.columns, rotation=30, ha="right")
    ax.set_yticks(range(len(rc.index)), rc.index)
    for i in range(rc.shape[0]):
        for j in range(rc.shape[1]):
            ax.text(j, i, f"{rc.values[i, j]:.1f}", ha="center", va="center", fontsize=9)
    ax.set_title("Net revenue by region and category (INR million)", loc="left", fontweight="bold")
    fig.tight_layout(); fig.savefig(DOCS / "region_category_heatmap.png", dpi=150); plt.close(fig)

    # before/after data quality
    rej_total = sum(metrics["orders_rejected_by_reason"].values())
    items = {
        "Duplicate orders removed": metrics["orders_duplicates_removed"],
        "Rows quarantined (invalid)": rej_total,
        "Missing prices filled": metrics["unit_price_missing_filled_from_master"],
        "Missing discounts filled": metrics["discount_missing_filled_with_0"],
        "Currency-formatted prices parsed": metrics["price_values_with_currency_symbols"],
    }
    fig, ax = plt.subplots(figsize=(9, 3.8))
    ax.barh(list(items)[::-1], list(items.values())[::-1], color="#C55A11")
    for y, v in enumerate(list(items.values())[::-1]):
        ax.text(v, y, f" {v:,}", va="center", fontsize=9)
    ax.set_title("Data quality issues found and fixed (orders)", loc="left", fontweight="bold")
    ax.set_xlim(0, max(items.values()) * 1.2)
    fig.tight_layout(); fig.savefig(DOCS / "data_quality_fixes.png", dpi=150); plt.close(fig)


# ---- Excel dashboard (formulas reference the Clean_Data sheet) ----
FONT = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BASE = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=16, bold=True, color="1F3864")
NOTE = Font(name=FONT, size=9, italic=True, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, row, col, labels):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_excel(flat, res, checks, path: Path):
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    dq = wb.create_sheet("Data_Quality")
    top = wb.create_sheet("Top_Customers_SQL")
    data = wb.create_sheet("Clean_Data")

    # ---------- Clean_Data (Excel Table, ready for a PivotTable) ----------
    cols = list(flat.columns)
    header(data, 1, 1, cols)
    for r, row in enumerate(flat.itertuples(index=False), start=2):
        for c, v in enumerate(row, start=1):
            data.cell(row=r, column=c, value=v)
    n = len(flat) + 1
    tbl = Table(displayName="SalesClean", ref=f"A1:{get_column_letter(len(cols))}{n}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    data.add_table(tbl)
    data.freeze_panes = "A2"
    for i, w in enumerate([12, 12, 9, 11, 12, 9, 11, 16, 9, 11, 9, 12, 11, 16], start=1):
        data.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, n + 1):
        data.cell(row=r, column=10).number_format = "#,##0.00"
        data.cell(row=r, column=11).number_format = "0%"
        data.cell(row=r, column=12).number_format = "#,##0.00"
    L = {name: get_column_letter(i + 1) for i, name in enumerate(cols)}

    def rng(name):
        return f"Clean_Data!${L[name]}$2:${L[name]}${n}"

    REV, STAT, MON, CAT, REG = rng("order_value"), rng("status"), rng("month"), rng("category"), rng("region")

    # ---------- Dashboard ----------
    dash["A1"], dash["A1"].font = "Sales Dashboard - Data Analyst Portfolio Project", TITLE
    dash["A2"] = ("Net revenue = order value of Delivered orders. All figures are live formulas on the "
                  "Clean_Data sheet. Dataset is synthetic.")
    dash["A2"].font = NOTE

    dash["A4"], dash["B4"] = "KPI", "Value"
    header(dash, 4, 1, ["KPI", "Value"])
    kpis = [
        ("Net revenue (INR)", f'=SUMIFS({REV},{STAT},"Delivered")', '#,##0'),
        ("Delivered orders", f'=COUNTIFS({STAT},"Delivered")', '#,##0'),
        ("Average order value (INR)", "=IFERROR(B5/B6,0)", '#,##0.00'),
        ("Total orders (all statuses)", f"=COUNTA({rng('order_id')})", '#,##0'),
        ("Cancellation rate", f'=IFERROR(COUNTIFS({STAT},"Cancelled")/B8,0)', '0.0%'),
        ("Return rate", f'=IFERROR(COUNTIFS({STAT},"Returned")/B8,0)', '0.0%'),
    ]
    for i, (lab, f, fmt) in enumerate(kpis, start=5):
        dash.cell(row=i, column=1, value=lab).font = BOLD
        c = dash.cell(row=i, column=2, value=f)
        c.number_format, c.font = fmt, BASE
        for col in (1, 2):
            dash.cell(row=i, column=col).border = BOX

    # monthly table
    months = sorted(flat["month"].unique())
    m0 = 13
    dash.cell(row=m0 - 1, column=1, value="Monthly net revenue").font = BOLD
    header(dash, m0, 1, ["Month", "Net revenue", "Delivered orders", "MoM growth"])
    for i, mo in enumerate(months, start=m0 + 1):
        dash.cell(row=i, column=1, value=mo)
        dash.cell(row=i, column=2, value=f'=SUMIFS({REV},{STAT},"Delivered",{MON},A{i})').number_format = "#,##0"
        dash.cell(row=i, column=3, value=f'=COUNTIFS({STAT},"Delivered",{MON},A{i})').number_format = "#,##0"
        if i > m0 + 1:
            dash.cell(row=i, column=4, value=f"=IFERROR(B{i}/B{i-1}-1,0)").number_format = "0.0%"
        for col in range(1, 5):
            dash.cell(row=i, column=col).border = BOX
            dash.cell(row=i, column=col).font = BASE
    m1 = m0 + len(months)

    # category table
    cats = sorted(flat["category"].unique())
    c0 = m1 + 3
    dash.cell(row=c0 - 1, column=1, value="Net revenue by category").font = BOLD
    header(dash, c0, 1, ["Category", "Net revenue", "Share"])
    for i, ca in enumerate(cats, start=c0 + 1):
        dash.cell(row=i, column=1, value=ca)
        dash.cell(row=i, column=2, value=f'=SUMIFS({REV},{STAT},"Delivered",{CAT},A{i})').number_format = "#,##0"
        dash.cell(row=i, column=3, value=f"=IFERROR(B{i}/$B$5,0)").number_format = "0.0%"
        for col in range(1, 4):
            dash.cell(row=i, column=col).border = BOX
            dash.cell(row=i, column=col).font = BASE
    c1 = c0 + len(cats)

    # region x category matrix
    regs = sorted(flat["region"].unique())
    x0 = c1 + 3
    dash.cell(row=x0 - 1, column=1, value="Net revenue: region x category").font = BOLD
    header(dash, x0, 1, ["Region"] + cats)
    for i, rg in enumerate(regs, start=x0 + 1):
        dash.cell(row=i, column=1, value=rg).font = BOLD
        dash.cell(row=i, column=1).border = BOX
        for j, ca in enumerate(cats, start=2):
            col = get_column_letter(j)
            cell = dash.cell(row=i, column=j,
                             value=f'=SUMIFS({REV},{STAT},"Delivered",{REG},$A{i},{CAT},{col}${x0})')
            cell.number_format, cell.font, cell.border = "#,##0", BASE, BOX
    x1 = x0 + len(regs)
    dash.conditional_formatting.add(
        f"B{x0+1}:{get_column_letter(len(cats)+1)}{x1}",
        ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="2E75B6"))

    note_r = x1 + 2
    dash.cell(row=note_r, column=1, value="Add a PivotTable:").font = BOLD
    dash.cell(row=note_r + 1, column=1,
              value="Clean_Data sheet > click any cell in the table > Insert > PivotTable > new sheet. "
                    "Try Rows = category, Columns = region, Values = Sum of order_value, Filter = status.").font = NOTE

    # charts
    lc = LineChart(); lc.title = "Monthly net revenue"; lc.height, lc.width = 7.5, 16
    lc.add_data(Reference(dash, min_col=2, min_row=m0, max_row=m1), titles_from_data=True)
    lc.set_categories(Reference(dash, min_col=1, min_row=m0 + 1, max_row=m1))
    lc.legend = None
    dash.add_chart(lc, "F4")
    bc = BarChart(); bc.type = "bar"; bc.title = "Net revenue by category"; bc.height, bc.width = 7.5, 16
    bc.add_data(Reference(dash, min_col=2, min_row=c0, max_row=c1), titles_from_data=True)
    bc.set_categories(Reference(dash, min_col=1, min_row=c0 + 1, max_row=c1))
    bc.legend = None
    dash.add_chart(bc, "F20")

    dash.column_dimensions["A"].width = 30
    for col in "BCDEFGH":
        dash.column_dimensions[col].width = 16
    dash.sheet_view.showGridLines = False

    # ---------- Data_Quality ----------
    dq["A1"], dq["A1"].font = "Data Quality Summary", TITLE
    dq["A2"] = "Values below are logged by run_pipeline.py when the raw data was cleaned (hardcoded results of the pipeline run)."
    dq["A2"].font = NOTE
    header(dq, 4, 1, ["Issue", "Count", "Action taken"])
    m = metrics
    rows = [
        ("Duplicate order rows", m["orders_duplicates_removed"], "Removed, kept first"),
        ("Duplicate customer rows", m["customers_duplicates_removed"], "Removed, kept first"),
        ("Price values with currency symbols", m["price_values_with_currency_symbols"], "Parsed to numbers"),
        ("Missing unit prices", m["unit_price_missing_filled_from_master"], "Filled from product master"),
        ("Missing discounts", m["discount_missing_filled_with_0"], "Assumed 0%"),
        ("Invalid customer emails", m["emails_invalid"], "Set to blank"),
        ("Missing customer emails", m["emails_missing"], "Left blank"),
    ] + [(f"Quarantined: {k}", v, "Moved to rejected file") for k, v in m["orders_rejected_by_reason"].items()]
    for i, (a, b, c) in enumerate(rows, start=5):
        for j, v in enumerate((a, b, c), start=1):
            cell = dq.cell(row=i, column=j, value=v)
            cell.font, cell.border = BASE, BOX
        dq.cell(row=i, column=2).number_format = "#,##0"
    r = 5 + len(rows) + 2
    dq.cell(row=r, column=1, value="Post-cleaning validation").font = BOLD
    header(dq, r + 1, 1, ["Check", "Result"])
    for i, ch in enumerate(checks, start=r + 2):
        dq.cell(row=i, column=1, value=ch["check"]).font = BASE
        res_c = dq.cell(row=i, column=2, value=ch["result"])
        res_c.font = Font(name=FONT, size=10, bold=True, color="006100" if ch["result"] == "PASS" else "C00000")
        for j in (1, 2):
            dq.cell(row=i, column=j).border = BOX
    dq.column_dimensions["A"].width = 48; dq.column_dimensions["B"].width = 12; dq.column_dimensions["C"].width = 30

    # ---------- Top_Customers_SQL ----------
    top["A1"], top["A1"].font = "Top 10 customers by net revenue", TITLE
    top["A2"] = "Source: SQL query q05_top10_customers (RANK window function) on db/sales.db. Static values."
    top["A2"].font = NOTE
    t = res["q05_top10_customers"]
    header(top, 4, 1, list(t.columns))
    for i, row in enumerate(t.itertuples(index=False), start=5):
        for j, v in enumerate(row, start=1):
            cell = top.cell(row=i, column=j, value=v)
            cell.font, cell.border = BASE, BOX
        top.cell(row=i, column=3).number_format = "#,##0.00"
    for col, w in zip("ABCDE", (14, 14, 16, 10, 14)):
        top.column_dimensions[col].width = w

    wb.save(path)


def recalc(path: Path):
    script = "/mnt/skills/public/xlsx/scripts/recalc.py"
    if Path(script).exists():
        out = subprocess.run([sys.executable, script, str(path), "120"], capture_output=True, text=True)
        print("recalc:", out.stdout.strip()[:400])


# ------------------------------------------------------------------ main
def main():
    if not (RAW / "orders_raw.csv").exists():
        import generate_data
        generate_data.main()

    print("1/6 Extract")
    orders_raw, customers_raw, products_raw = extract()

    print("2/6 Clean")
    products = clean_products(products_raw, orders_raw)
    customers = clean_customers(customers_raw)
    orders, rejected = clean_orders(orders_raw, customers, products)
    metrics["clean_rows"] = {"orders": len(orders), "customers": len(customers), "products": len(products)}

    print("3/6 Validate")
    checks = validate(orders, customers, products)
    failed = [c for c in checks if c["result"] == "FAIL"]
    if failed:
        print("VALIDATION FAILED:", failed)
        sys.exit(1)

    CLEAN.mkdir(exist_ok=True); REJ.mkdir(exist_ok=True); REPORTS.mkdir(exist_ok=True)
    orders.to_csv(CLEAN / "orders_clean.csv", index=False)
    customers.to_csv(CLEAN / "customers_clean.csv", index=False)
    products.to_csv(CLEAN / "products_clean.csv", index=False)
    rejected.to_csv(REJ / "orders_rejected.csv", index=False)

    print("4/6 Load to SQLite")
    con = load_sqlite(orders, customers, products)

    print("5/6 Run SQL queries")
    res = run_sql(con)
    flat = make_flat(con)
    con.close()

    print("6/6 Reports")
    write_quality_report(checks)
    make_charts(res, flat)
    xlsx = REPORTS / "sales_dashboard.xlsx"
    build_excel(flat, res, checks, xlsx)
    recalc(xlsx)

    rej_total = sum(metrics["orders_rejected_by_reason"].values())
    print(f"\nDone. Raw orders {metrics['raw_rows']['orders']:,} -> clean {len(orders):,} "
          f"(duplicates {metrics['orders_duplicates_removed']:,}, quarantined {rej_total:,}).")


if __name__ == "__main__":
    main()
